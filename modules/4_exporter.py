"""
Module 4: Master BOQ & Global Material Excel Exporter
Generates executive-ready, AEC-standard Master Bill of Quantities (BOQ) spreadsheets (.xlsx)
using openpyxl, featuring:
- Sheet 1: Master Summary (Global KPIs, Material Grouped Totals, Page-by-Page breakdown)
- Sheet 2: Detailed BOQ (Page No., Drawing/Room Name, Material Type, Panel ID, L x W, Cutouts, Net Area)
- Sheet 3: Nesting by Material (Independent 2D nesting BOM, cutting layouts, CNC coordinates grouped by material)
- Sheet 4: QC Audit Log (if QC report provided)
"""

from datetime import datetime
import io
import math
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

import importlib

_mod_geom = importlib.import_module("modules.2_geometry")
PanelItem = _mod_geom.PanelItem

_mod_nest = importlib.import_module("modules.3_nesting")
NestingResult = _mod_nest.NestingResult
GlobalNestingResult = _mod_nest.GlobalNestingResult
NestedSheet = _mod_nest.NestedSheet


def add_api_cost_report_sheet(
    wb: Any,
    api_usage_logs: Optional[List[Dict[str, Any]]] = None,
    project_name: str = "Acoustic Takeoff Master Project"
) -> Any:
    """
    Generate an AEC-standard 'API Cost Report' sheet within the Master BOQ Excel Workbook.
    
    Features:
    - Tracks AI resource consumption:
        * Gemini 1.5 Flash (or Pro/Vision) Token usage ($0.075 / 1M tokens)
        * Replicate SAM 2 GPU execution seconds ($0.000725 / sec)
    - Executive KPI Cards: Total API Cost, Total Gemini Tokens, SAM 2 GPU Time, Total Requests.
    - Standard tabular breakdown: Thời gian, Loại API, Lượng sử dụng, Đơn vị tính, Đơn giá tham khảo, Thành tiền (USD), Diễn giải tác vụ.
    - Dynamic formula total row: =SUM(G...:G...) formatted as Currency ($).
    - Exception handling: If api_usage_logs is empty (e.g. Native Vector Engine mode),
      displays prominent message 'Không phát sinh chi phí AI (Sử dụng Vector Engine)' with $0.00 cost.
    """
    if openpyxl is None:
        return None

    is_bytes_input = isinstance(wb, (bytes, bytearray))
    if is_bytes_input:
        wb_obj = openpyxl.load_workbook(io.BytesIO(wb))
    else:
        wb_obj = wb

    FONT_FAMILY = "Segoe UI"
    font_main_title = Font(name=FONT_FAMILY, size=15, bold=True, color="1E293B")
    font_subtitle = Font(name=FONT_FAMILY, size=8.5, color="64748B", italic=True)
    font_section_hdr = Font(name=FONT_FAMILY, size=11, bold=True, color="0F172A")
    font_table_hdr = Font(name=FONT_FAMILY, size=9.5, bold=True, color="FFFFFF")
    font_bold = Font(name=FONT_FAMILY, size=9.5, bold=True, color="1E293B")
    font_normal = Font(name=FONT_FAMILY, size=9, color="334155")
    font_kpi_value = Font(name=FONT_FAMILY, size=15, bold=True, color="0369A1")
    font_kpi_label = Font(name=FONT_FAMILY, size=8, bold=True, color="64748B")
    font_notice_title = Font(name=FONT_FAMILY, size=11, bold=True, color="0369A1")
    font_notice_sub = Font(name=FONT_FAMILY, size=9, color="475569", italic=True)

    fill_navy_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue_header = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    fill_slate_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_total_row = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_kpi_card = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")
    fill_notice = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    thin_border = Side(border_style="thin", color="CBD5E1")
    double_bottom = Side(border_style="double", color="1E293B")
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    total_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=double_bottom)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Remove existing sheet if replacing
    if "API Cost Report" in wb_obj.sheetnames:
        del wb_obj["API Cost Report"]

    ws = wb_obj.create_sheet(title="API Cost Report")
    ws.views.sheetView[0].showGridLines = True

    # Title Block
    ws["A2"] = "BÁO CÁO CHI PHÍ DỊCH VỤ API & TÀI NGUYÊN AI (API COST REPORT)"
    ws["A2"].font = font_main_title
    ws["A3"] = f"Dự án: {project_name} | Hệ Thống Bóc Tách Khối Lượng Tự Động AEC & Multi-Agent Vision | Ngày tạo: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A3"].font = font_subtitle

    # Normalize empty or None logs
    logs = list(api_usage_logs) if api_usage_logs else []

    # CASE 1: No API usage (Native Vector Engine mode)
    if not logs:
        # KPI Cards (A5:D6)
        ws["A5"] = "TỔNG CHI PHÍ PHÁT SINH"
        ws["A5"].font = font_kpi_label
        ws["A5"].fill = fill_kpi_card
        ws["A5"].border = cell_border
        ws["A6"] = 0.0
        ws["A6"].number_format = '"$"#,##0.0000'
        ws["A6"].font = font_kpi_value
        ws["A6"].fill = fill_kpi_card
        ws["A6"].border = cell_border
        ws["A6"].alignment = align_center

        ws["B5"] = "ĐỘNG CƠ BÓC TÁCH"
        ws["B5"].font = font_kpi_label
        ws["B5"].fill = fill_kpi_card
        ws["B5"].border = cell_border
        ws["B6"] = "Vector Engine (Offline)"
        ws["B6"].font = font_bold
        ws["B6"].fill = fill_kpi_card
        ws["B6"].border = cell_border
        ws["B6"].alignment = align_center

        ws["C5"] = "SỐ PHIÊN GỌI CLOUD API"
        ws["C5"].font = font_kpi_label
        ws["C5"].fill = fill_kpi_card
        ws["C5"].border = cell_border
        ws["C6"] = 0
        ws["C6"].font = font_bold
        ws["C6"].fill = fill_kpi_card
        ws["C6"].border = cell_border
        ws["C6"].alignment = align_center

        ws["D5"] = "TIÊU THỤ TÀI NGUYÊN"
        ws["D5"].font = font_kpi_label
        ws["D5"].fill = fill_kpi_card
        ws["D5"].border = cell_border
        ws["D6"] = "0 Tokens / 0.0s GPU"
        ws["D6"].font = font_bold
        ws["D6"].fill = fill_kpi_card
        ws["D6"].border = cell_border
        ws["D6"].alignment = align_center

        # Notice Banner (Requirement 3: "Không phát sinh chi phí AI (Sử dụng Vector Engine)")
        ws["A8"] = "TRẠNG THÁI TÀI NGUYÊN HỆ THỐNG"
        ws["A8"].font = font_section_hdr

        ws["A9"] = "Không phát sinh chi phí AI (Sử dụng Vector Engine)"
        ws["A9"].font = font_notice_title
        ws["A9"].fill = fill_notice

        ws["A10"] = "Toàn bộ khối lượng bóc tách được xử lý hoàn toàn bằng Động cơ Vector Bản địa (PyMuPDF & Shapely) chạy offline trên máy trạm. Không tiêu tốn Token Gemini LLM hay thời gian tính toán GPU Replicate SAM 2."
        ws["A10"].font = font_notice_sub

        for col_idx in range(1, 9):
            ws.cell(row=9, column=col_idx).fill = fill_notice
            ws.cell(row=9, column=col_idx).border = cell_border

        # Table
        start_row = 12
        headers = ["STT", "Thời gian", "Loại API", "Lượng sử dụng (Tokens / Seconds)", "Đơn vị tính", "Đơn giá tham khảo (USD)", "Thành tiền (USD)", "Diễn giải tác vụ"]
        for c_idx, h in enumerate(headers, start=1):
            c_cell = ws.cell(row=start_row, column=c_idx, value=h)
            c_cell.font = font_table_hdr
            c_cell.fill = fill_navy_header
            c_cell.alignment = align_center
            c_cell.border = cell_border

        # Informative Row
        ws.cell(row=13, column=1, value=1).alignment = align_center
        ws.cell(row=13, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S")).alignment = align_center
        ws.cell(row=13, column=3, value="Native Vector Engine (PyMuPDF & Shapely)").font = font_bold
        ws.cell(row=13, column=4, value=0).alignment = align_right
        ws.cell(row=13, column=5, value="Local Ops").alignment = align_center
        ws.cell(row=13, column=6, value="$0.00 / Local CPU").alignment = align_center
        c_amt = ws.cell(row=13, column=7, value=0.0)
        c_amt.number_format = '"$"#,##0.0000'
        c_amt.alignment = align_right
        c_amt.font = font_bold
        ws.cell(row=13, column=8, value="Bóc tách vector 2D nội bộ - Không phát sinh chi phí AI đám mây.").font = font_normal

        for c_idx in range(1, 9):
            ws.cell(row=13, column=c_idx).border = cell_border

        # Total Row
        ws.cell(row=14, column=1, value="TỔNG CỘNG CHI PHÍ API (TOTAL COST)").font = font_bold
        ws.cell(row=14, column=1).alignment = align_left
        tot_c = ws.cell(row=14, column=7, value=0.0)
        tot_c.number_format = '"$"#,##0.0000'
        tot_c.font = font_bold
        tot_c.alignment = align_right
        ws.cell(row=14, column=8, value="Tổng chi phí API = $0.00 USD").alignment = align_left
        ws.cell(row=14, column=8).font = font_subtitle

        for c_idx in range(1, 9):
            ws.cell(row=14, column=c_idx).border = total_border
            ws.cell(row=14, column=c_idx).fill = fill_total_row

        return ws

    # CASE 2: Active API usage logs
    tot_tokens = 0
    tot_sam2_secs = 0.0
    tot_cost = 0.0

    processed_logs = []
    for entry in logs:
        api_type = str(entry.get("api_type", entry.get("type", "Gemini 1.5 Flash")))
        unit = str(entry.get("unit", "Tokens"))
        usage = float(entry.get("usage_amount", entry.get("usage", 0.0)))
        ts = str(entry.get("timestamp", entry.get("time", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))))
        desc = str(entry.get("description", entry.get("notes", "-")))

        api_lower = api_type.lower()
        if "token" in unit.lower() or "gemini" in api_lower:
            tot_tokens += int(usage)
            if "pro" in api_lower:
                rate_val = 3.50 / 1_000_000
                rate_label = "$3.50 / 1M tokens"
            else: # Gemini 1.5 Flash default: $0.075 / 1M tokens
                rate_val = 0.075 / 1_000_000
                rate_label = "$0.075 / 1M tokens"
            calc_cost = usage * rate_val
        elif "sec" in unit.lower() or "sam" in api_lower or "replicate" in api_lower:
            tot_sam2_secs += usage
            # Replicate SAM 2 default: $0.000725 / sec
            rate_val = 0.000725
            rate_label = "$0.000725 / giây"
            calc_cost = usage * rate_val
        else:
            rate_val = float(entry.get("unit_rate", 0.0))
            rate_label = str(entry.get("rate_label", f"${rate_val:.6f}"))
            calc_cost = usage * rate_val

        cost_usd = float(entry.get("cost_usd", entry.get("cost", calc_cost)))
        tot_cost += cost_usd

        processed_logs.append({
            "timestamp": ts,
            "api_type": api_type,
            "usage": usage,
            "unit": unit,
            "rate_label": rate_label,
            "cost_usd": cost_usd,
            "description": desc
        })

    # KPI Cards (A5:D6)
    ws["A5"] = "TỔNG CHI PHÍ API (USD)"
    ws["A5"].font = font_kpi_label
    ws["A5"].fill = fill_kpi_card
    ws["A5"].border = cell_border
    ws["A6"] = round(tot_cost, 4)
    ws["A6"].number_format = '"$"#,##0.0000'
    ws["A6"].font = font_kpi_value
    ws["A6"].fill = fill_kpi_card
    ws["A6"].border = cell_border
    ws["A6"].alignment = align_center

    ws["B5"] = "TỔNG TOKENS GEMINI"
    ws["B5"].font = font_kpi_label
    ws["B5"].fill = fill_kpi_card
    ws["B5"].border = cell_border
    ws["B6"] = f"{tot_tokens:,} Tokens"
    ws["B6"].font = font_bold
    ws["B6"].fill = fill_kpi_card
    ws["B6"].border = cell_border
    ws["B6"].alignment = align_center

    ws["C5"] = "THỜI GIAN GPU SAM 2"
    ws["C5"].font = font_kpi_label
    ws["C5"].fill = fill_kpi_card
    ws["C5"].border = cell_border
    ws["C6"] = f"{tot_sam2_secs:.2f} Giây"
    ws["C6"].font = font_bold
    ws["C6"].fill = fill_kpi_card
    ws["C6"].border = cell_border
    ws["C6"].alignment = align_center

    ws["D5"] = "SỐ PHIÊN GỌI DỊCH VỤ"
    ws["D5"].font = font_kpi_label
    ws["D5"].fill = fill_kpi_card
    ws["D5"].border = cell_border
    ws["D6"] = f"{len(processed_logs)} Phiên"
    ws["D6"].font = font_bold
    ws["D6"].fill = fill_kpi_card
    ws["D6"].border = cell_border
    ws["D6"].alignment = align_center

    ws["A8"] = "CHI TIẾT LỊCH SỬ TIÊU THỤ TÀI NGUYÊN ĐÁM MÂY (API CONSUMPTION BREAKDOWN)"
    ws["A8"].font = font_section_hdr

    # Table Headers
    start_row = 10
    headers = ["STT", "Thời gian", "Loại API", "Lượng sử dụng (Tokens / Seconds)", "Đơn vị tính", "Đơn giá tham khảo (USD)", "Thành tiền (USD)", "Diễn giải tác vụ"]
    for c_idx, h in enumerate(headers, start=1):
        c_cell = ws.cell(row=start_row, column=c_idx, value=h)
        c_cell.font = font_table_hdr
        c_cell.fill = fill_navy_header
        c_cell.alignment = align_center
        c_cell.border = cell_border

    # Data Rows
    cur_row = start_row + 1
    for idx, item in enumerate(processed_logs, start=1):
        ws.cell(row=cur_row, column=1, value=idx).alignment = align_center
        ws.cell(row=cur_row, column=2, value=item["timestamp"]).alignment = align_center
        ws.cell(row=cur_row, column=3, value=item["api_type"]).font = font_bold

        usage_cell = ws.cell(row=cur_row, column=4, value=item["usage"])
        if "token" in item["unit"].lower():
            usage_cell.number_format = '#,##0'
        else:
            usage_cell.number_format = '#,##0.00'
        usage_cell.alignment = align_right

        ws.cell(row=cur_row, column=5, value=item["unit"]).alignment = align_center
        ws.cell(row=cur_row, column=6, value=item["rate_label"]).alignment = align_center

        cost_cell = ws.cell(row=cur_row, column=7, value=item["cost_usd"])
        cost_cell.number_format = '"$"#,##0.0000'
        cost_cell.font = font_bold
        cost_cell.alignment = align_right

        ws.cell(row=cur_row, column=8, value=item["description"]).font = font_normal

        for c_idx in range(1, 9):
            c_c = ws.cell(row=cur_row, column=c_idx)
            c_c.border = cell_border
            if idx % 2 == 1:
                c_c.fill = fill_zebra
        cur_row += 1

    # Total Row
    ws.cell(row=cur_row, column=1, value="TỔNG CỘNG CHI PHÍ API (TOTAL ESTIMATED COST)").font = font_bold
    ws.cell(row=cur_row, column=1).alignment = align_left

    tot_cell = ws.cell(row=cur_row, column=7, value=f"=SUM(G{start_row + 1}:G{cur_row - 1})")
    tot_cell.font = font_bold
    tot_cell.number_format = '"$"#,##0.0000'
    tot_cell.alignment = align_right

    ws.cell(row=cur_row, column=8, value="Tổng chi phí điện toán đám mây cho phiên bóc tách.").alignment = align_left
    ws.cell(row=cur_row, column=8).font = font_subtitle

    for c_idx in range(1, 9):
        c_cell = ws.cell(row=cur_row, column=c_idx)
        c_cell.border = total_border
        c_cell.fill = fill_total_row

    if is_bytes_input:
        buf = io.BytesIO()
        wb_obj.save(buf)
        return buf.getvalue()
    return ws


def export_takeoff_to_excel(
    panels: List[PanelItem],
    nesting: Optional[Union[GlobalNestingResult, NestingResult]] = None,
    qc_report: Optional[Any] = None,
    project_name: str = "Acoustic Takeoff Master Project",
    drawing_ref: str = "Full Project Takeoff",
    scale_label: str = "1:50",
    cost_per_sheet: float = 65.0,
    cost_per_m2: float = 22.5,
    unit: str = "mm",
    output_path: Optional[Union[str, Path]] = None,
    material_cost_rates: Optional[Dict[str, float]] = None,
    api_usage_logs: Optional[List[Dict[str, Any]]] = None
) -> bytes:
    """
    Generate an AEC-standard Master BOQ Excel Workbook (.xlsx):
    - Sheet 1: Master Summary (Global KPIs, Material-Grouped Summary, Page Breakdown)
    - Sheet 2: Detailed BOQ (Itemized panel schedule with Page No., Drawing/Room Name, Material Type, Panel ID, L x W, Cutouts, Net Area)
    - Sheet 3: Nesting by Material (Nesting BOM, cutting plans & CNC coordinates separated by material)
    - Sheet 4: Quality Control & Drafting Audit Log (optional)
    - Sheet 5: API Cost Report (Cloud AI Token & GPU usage, cost breakdown)
    """
    if openpyxl is None:
        raise ImportError("openpyxl is required for Excel export. Please install it with `pip install openpyxl`.")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # Professional AEC Styling Palette (Navy, Slate, Steel, Pure White)
    FONT_FAMILY = "Segoe UI"
    font_main_title = Font(name=FONT_FAMILY, size=15, bold=True, color="1E293B")
    font_subtitle = Font(name=FONT_FAMILY, size=8.5, color="64748B", italic=True)
    font_section_hdr = Font(name=FONT_FAMILY, size=11, bold=True, color="0F172A")
    font_table_hdr = Font(name=FONT_FAMILY, size=9.5, bold=True, color="FFFFFF")
    font_mat_hdr = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    font_bold = Font(name=FONT_FAMILY, size=9.5, bold=True, color="1E293B")
    font_normal = Font(name=FONT_FAMILY, size=9, color="334155")
    font_kpi_value = Font(name=FONT_FAMILY, size=16, bold=True, color="0369A1")
    font_kpi_label = Font(name=FONT_FAMILY, size=8, bold=True, color="64748B")

    # Fills
    fill_navy_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_blue_header = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    fill_slate_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_amber_header = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_total_row = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_kpi_card = PatternFill(start_color="F0F9FF", end_color="F0F9FF", fill_type="solid")

    # Borders
    thin_border = Side(border_style="thin", color="CBD5E1")
    thick_bottom = Side(border_style="medium", color="1E293B")
    double_bottom = Side(border_style="double", color="1E293B")
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    total_border = Border(top=thin_border, bottom=double_bottom)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Rates lookup
    rates = material_cost_rates or {
        "Polyester 9mm": cost_per_sheet,
        "Wood Slat Panel": cost_per_sheet * 1.6,
        "Fabric Wrapped 25mm": cost_per_sheet * 1.35
    }

    # Normalize Nesting Results
    material_nest_dict: Dict[str, NestingResult] = {}
    if isinstance(nesting, GlobalNestingResult):
        material_nest_dict = nesting.material_results
    elif isinstance(nesting, NestingResult):
        material_nest_dict = {nesting.material_type: nesting}

    # Calculate Global Project Metrics
    total_panels = len(panels)
    total_net_area = round(sum(p.net_area for p in panels), 3)
    total_gross_area = round(sum(p.gross_area for p in panels), 3)
    total_cutout_area = round(sum(p.cutout_area for p in panels), 3)

    total_sheets = sum(res.total_sheets for res in material_nest_dict.values()) if material_nest_dict else (nesting.total_sheets if nesting else 0)
    total_stock_area = sum(res.total_stock_area_m2 for res in material_nest_dict.values()) if material_nest_dict else 0.0
    total_used_area = sum(res.total_used_area_m2 for res in material_nest_dict.values()) if material_nest_dict else 0.0
    overall_yield = round((total_used_area / max(0.001, total_stock_area)) * 100.0, 1) if total_stock_area > 0 else 0.0
    overall_waste = round(100.0 - overall_yield, 1) if total_stock_area > 0 else 0.0

    total_est_cost = sum(
        res.total_sheets * rates.get(m_name, cost_per_sheet)
        for m_name, res in material_nest_dict.items()
    ) if material_nest_dict else total_sheets * cost_per_sheet

    # Group Panels by Material
    mat_groups: Dict[str, List[PanelItem]] = {}
    for p in panels:
        mat_groups.setdefault(p.material_type, []).append(p)

    # Group Panels by Page
    page_groups: Dict[int, List[PanelItem]] = {}
    for p in panels:
        page_groups.setdefault(p.page_number, []).append(p)

    # =========================================================================
    # SHEET 1: MASTER SUMMARY
    # =========================================================================
    ws1 = wb.create_sheet(title="Master Summary")
    ws1.views.sheetView[0].showGridLines = True

    # Title Block
    ws1["B2"] = "MASTER BILL OF QUANTITIES (BOQ) & ACOUSTIC TAKEOFF SUMMARY"
    ws1["B2"].font = font_main_title
    ws1["B3"] = f"Generated by Antigravity AEC Batch Takeoff & Global Nesting Engine | {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["B3"].font = font_subtitle

    # Project Information Card
    ws1["B5"] = "PROJECT SPECIFICATIONS"
    ws1["B5"].font = font_section_hdr
    meta_info = [
        ("Project Name:", project_name),
        ("Document / Package:", drawing_ref),
        ("Total Pages Processed:", f"{len(page_groups)} Drawing Pages with Acoustic Scope"),
        ("Drawing Scale Preset:", scale_label),
        ("Fabrication Units:", unit.upper()),
        ("Report Export Date:", datetime.now().strftime("%Y-%m-%d")),
        ("Global Nesting Scope:", "Independent Project-Wide Material Optimization")
    ]
    for idx, (label, val) in enumerate(meta_info, start=6):
        ws1[f"B{idx}"] = label
        ws1[f"B{idx}"].font = font_bold
        ws1[f"C{idx}"] = val
        ws1[f"C{idx}"].font = font_normal

    # KPI Summary Cards in Columns E, G, I
    kpi_cards = [
        ("TOTAL NET ACOUSTIC AREA", f"{total_net_area:,.2f} m²", "E5:F7"),
        ("TOTAL PANELS / TILES", f"{total_panels:,} pcs", "G5:H7"),
        ("TOTAL STOCK SHEETS REQ.", f"{total_sheets:,} sheets", "I5:J7"),
        ("OVERALL MATERIAL YIELD", f"{overall_yield:.1f}%", "E9:F11"),
        ("SCRAP / OFFCUT WASTE", f"{overall_waste:.1f}%", "G9:H11"),
        ("ESTIMATED SHEET COST", f"${total_est_cost:,.2f}", "I9:J11")
    ]
    for title, val, rng in kpi_cards:
        ws1.merge_cells(rng)
        top_left = rng.split(":")[0]
        cell = ws1[top_left]
        cell.value = f"{title}\n{val}"
        cell.alignment = align_center
        cell.fill = fill_kpi_card
        cell.font = font_kpi_value

        # Apply borders across merged card
        col_s, row_s = top_left[0], int(top_left[1:])
        col_e, row_e = rng.split(":")[1][0], int(rng.split(":")[1][1:])
        for r in range(row_s, row_e + 1):
            for c in range(ord(col_s) - ord('A') + 1, ord(col_e) - ord('A') + 2):
                ws1.cell(row=r, column=c).border = cell_border

    # Master Table 1: Totals Grouped by Material Type
    ws1["B14"] = "1. ACOUSTIC TAKEOFF & STOCK REQUIREMENTS (GROUPED BY MATERIAL TYPE)"
    ws1["B14"].font = font_section_hdr

    mat_table_headers = [
        "Material Specification", "Piece Count", "Net Area (m²)", "Stock Sheets Req.",
        "Total Stock Area (m²)", "Yield %", "Waste %", "Sheet Rate ($)", "Est. Material Cost ($)"
    ]
    for c_idx, h in enumerate(mat_table_headers, start=2):
        cell = ws1.cell(row=16, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_navy_header
        cell.alignment = align_center

    curr_row = 17
    for mat_name, p_list in sorted(mat_groups.items()):
        mat_pcs = len(p_list)
        mat_net = sum(p.net_area for p in p_list)
        res = material_nest_dict.get(mat_name)
        mat_sheets = res.total_sheets if res else math.ceil(mat_net / 2.88)
        mat_stk_area = res.total_stock_area_m2 if res else mat_sheets * 2.88
        mat_yield = (mat_net / max(0.001, mat_stk_area)) if mat_stk_area > 0 else 0.0
        mat_waste = max(0.0, 1.0 - mat_yield)
        rate = rates.get(mat_name, cost_per_sheet)
        mat_cost = mat_sheets * rate

        ws1.cell(row=curr_row, column=2, value=mat_name).font = font_bold
        ws1.cell(row=curr_row, column=3, value=mat_pcs).alignment = align_right
        ws1.cell(row=curr_row, column=4, value=round(mat_net, 3)).number_format = "#,##0.00"
        ws1.cell(row=curr_row, column=5, value=mat_sheets).alignment = align_right
        ws1.cell(row=curr_row, column=6, value=round(mat_stk_area, 2)).number_format = "#,##0.00"
        ws1.cell(row=curr_row, column=7, value=mat_yield).number_format = "0.0%"
        ws1.cell(row=curr_row, column=8, value=mat_waste).number_format = "0.0%"
        ws1.cell(row=curr_row, column=9, value=rate).number_format = "$#,##0.00"
        ws1.cell(row=curr_row, column=10, value=mat_cost).number_format = "$#,##0.00"

        for col in range(2, 11):
            c = ws1.cell(row=curr_row, column=col)
            c.border = cell_border
            if curr_row % 2 == 1:
                c.fill = fill_zebra
        curr_row += 1

    # Totals Row for Material Table
    ws1.cell(row=curr_row, column=2, value="TOTAL PROJECT SCOPE").font = font_bold
    ws1.cell(row=curr_row, column=3, value=f"=SUM(C17:C{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=4, value=f"=SUM(D17:D{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=4).number_format = "#,##0.00"
    ws1.cell(row=curr_row, column=5, value=f"=SUM(E17:E{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=6, value=f"=SUM(F17:F{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=6).number_format = "#,##0.00"
    ws1.cell(row=curr_row, column=7, value=f"=D{curr_row}/F{curr_row}").font = font_bold
    ws1.cell(row=curr_row, column=7).number_format = "0.0%"
    ws1.cell(row=curr_row, column=8, value=f"=1-G{curr_row}").font = font_bold
    ws1.cell(row=curr_row, column=8).number_format = "0.0%"
    ws1.cell(row=curr_row, column=9, value="-").alignment = align_center
    ws1.cell(row=curr_row, column=10, value=f"=SUM(J17:J{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=10).number_format = "$#,##0.00"

    for col in range(2, 11):
        c = ws1.cell(row=curr_row, column=col)
        c.border = total_border
        c.fill = fill_total_row

    # Master Table 2: Page-by-Page Scope Breakdown
    curr_row += 3
    ws1.cell(row=curr_row, column=2, value="2. PAGE-BY-PAGE DRAWING TAKEOFF SUMMARY").font = font_section_hdr
    curr_row += 2

    page_table_headers = [
        "Page No.", "Drawing / Room Name", "Materials Present", "Panel Count",
        "Gross Area (m²)", "Cutout Area (m²)", "Net Area (m²)", "% of Project Net Area"
    ]
    for c_idx, h in enumerate(page_table_headers, start=2):
        cell = ws1.cell(row=curr_row, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_slate_header
        cell.alignment = align_center

    page_start_row = curr_row + 1
    curr_row += 1

    for p_no in sorted(page_groups.keys()):
        p_items = page_groups[p_no]
        sh_name = p_items[0].sheet_name if p_items else f"Sheet {p_no}"
        mats_str = ", ".join(sorted(list(set(it.material_type for it in p_items))))
        p_count = len(p_items)
        p_gross = sum(it.gross_area for it in p_items)
        p_cut = sum(it.cutout_area for it in p_items)
        p_net = sum(it.net_area for it in p_items)
        p_pct = p_net / max(0.001, total_net_area)

        ws1.cell(row=curr_row, column=2, value=f"Page {p_no}").alignment = align_center
        ws1.cell(row=curr_row, column=3, value=sh_name).font = font_bold
        ws1.cell(row=curr_row, column=4, value=mats_str).font = font_normal
        ws1.cell(row=curr_row, column=5, value=p_count).alignment = align_right
        ws1.cell(row=curr_row, column=6, value=round(p_gross, 3)).number_format = "#,##0.00"
        ws1.cell(row=curr_row, column=7, value=round(p_cut, 3)).number_format = "#,##0.00"
        ws1.cell(row=curr_row, column=8, value=round(p_net, 3)).number_format = "#,##0.00"
        ws1.cell(row=curr_row, column=9, value=p_pct).number_format = "0.0%"

        for col in range(2, 10):
            c = ws1.cell(row=curr_row, column=col)
            c.border = cell_border
            if curr_row % 2 == 1:
                c.fill = fill_zebra
        curr_row += 1

    # Totals Row for Page Breakdown Table
    ws1.cell(row=curr_row, column=2, value="TOTALS").font = font_bold
    ws1.cell(row=curr_row, column=3, value="-").alignment = align_center
    ws1.cell(row=curr_row, column=4, value="-").alignment = align_center
    ws1.cell(row=curr_row, column=5, value=f"=SUM(E{page_start_row}:E{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=6, value=f"=SUM(F{page_start_row}:F{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=6).number_format = "#,##0.00"
    ws1.cell(row=curr_row, column=7, value=f"=SUM(G{page_start_row}:G{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=7).number_format = "#,##0.00"
    ws1.cell(row=curr_row, column=8, value=f"=SUM(H{page_start_row}:H{curr_row-1})").font = font_bold
    ws1.cell(row=curr_row, column=8).number_format = "#,##0.00"
    ws1.cell(row=curr_row, column=9, value="100.0%").font = font_bold

    for col in range(2, 10):
        c = ws1.cell(row=curr_row, column=col)
        c.border = total_border
        c.fill = fill_total_row

    # =========================================================================
    # SHEET 2: DETAILED BOQ (ITEMIZED PANEL SCHEDULE)
    # Required Columns: Page No., Drawing/Room Name, Material Type, Panel ID, L x W, Cutouts, Net Area
    # =========================================================================
    ws2 = wb.create_sheet(title="Detailed BOQ")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "DETAILED ITEM-BY-ITEM ACOUSTIC TAKEOFF SCHEDULE (MASTER BOQ)"
    ws2["A1"].font = font_section_hdr

    detailed_headers = [
        "Page No.", "Drawing/Room Name", "Material Type", "Panel ID", "Cut Length (mm)",
        "Cut Width (mm)", "L × W Format", "Cutouts Count", "Cutout Area (m²)",
        "Gross Area (m²)", "Net Area (m²)", "Perimeter (mm)", "Category", "CAD Layer"
    ]
    for c_idx, h in enumerate(detailed_headers, start=1):
        cell = ws2.cell(row=3, column=c_idx, value=h)
        cell.font = font_table_hdr
        cell.fill = fill_navy_header
        cell.alignment = align_center

    row_idx = 4
    for p in panels:
        ws2.cell(row=row_idx, column=1, value=f"Page {p.page_number}").alignment = align_center
        ws2.cell(row=row_idx, column=2, value=p.sheet_name).font = font_normal
        ws2.cell(row=row_idx, column=3, value=p.material_type).font = font_bold
        ws2.cell(row=row_idx, column=4, value=p.id).font = font_bold
        ws2.cell(row=row_idx, column=5, value=p.cut_length).number_format = "#,##0.0"
        ws2.cell(row=row_idx, column=6, value=p.cut_width).number_format = "#,##0.0"
        ws2.cell(row=row_idx, column=7, value=p.dimension_str).alignment = align_center
        ws2.cell(row=row_idx, column=8, value=p.cutout_count).alignment = align_center
        ws2.cell(row=row_idx, column=9, value=p.cutout_area).number_format = "#,##0.000"
        ws2.cell(row=row_idx, column=10, value=p.gross_area).number_format = "#,##0.000"
        ws2.cell(row=row_idx, column=11, value=p.net_area).number_format = "#,##0.000"
        ws2.cell(row=row_idx, column=12, value=p.perimeter).number_format = "#,##0.0"
        ws2.cell(row=row_idx, column=13, value=p.category).font = font_normal
        ws2.cell(row=row_idx, column=14, value=p.layer or "Default").font = font_normal

        for c_idx in range(1, 15):
            c_cell = ws2.cell(row=row_idx, column=c_idx)
            c_cell.border = cell_border
            if row_idx % 2 == 1:
                c_cell.fill = fill_zebra
        row_idx += 1

    # Detailed Schedule Totals Row
    ws2.cell(row=row_idx, column=1, value="TOTALS").font = font_bold
    ws2.cell(row=row_idx, column=2, value="-").alignment = align_center
    ws2.cell(row=row_idx, column=3, value="-").alignment = align_center
    ws2.cell(row=row_idx, column=4, value=f"{len(panels)} Panels").font = font_bold
    ws2.cell(row=row_idx, column=5, value="-").alignment = align_center
    ws2.cell(row=row_idx, column=6, value="-").alignment = align_center
    ws2.cell(row=row_idx, column=7, value="-").alignment = align_center
    ws2.cell(row=row_idx, column=8, value=f"=SUM(H4:H{row_idx-1})").font = font_bold
    ws2.cell(row=row_idx, column=9, value=f"=SUM(I4:I{row_idx-1})").font = font_bold
    ws2.cell(row=row_idx, column=9).number_format = "#,##0.000"
    ws2.cell(row=row_idx, column=10, value=f"=SUM(J4:J{row_idx-1})").font = font_bold
    ws2.cell(row=row_idx, column=10).number_format = "#,##0.000"
    ws2.cell(row=row_idx, column=11, value=f"=SUM(K4:K{row_idx-1})").font = font_bold
    ws2.cell(row=row_idx, column=11).number_format = "#,##0.000"
    ws2.cell(row=row_idx, column=12, value=f"=SUM(L4:L{row_idx-1})").font = font_bold
    ws2.cell(row=row_idx, column=12).number_format = "#,##0.0"
    ws2.cell(row=row_idx, column=13, value="-").alignment = align_center
    ws2.cell(row=row_idx, column=14, value="-").alignment = align_center

    for c_idx in range(1, 15):
        c_cell = ws2.cell(row=row_idx, column=c_idx)
        c_cell.border = total_border
        c_cell.fill = fill_total_row

    # =========================================================================
    # SHEET 3: NESTING BY MATERIAL (GLOBAL CUT PLANS SEPARATED BY MATERIAL)
    # =========================================================================
    ws3 = wb.create_sheet(title="Nesting by Material")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = "STOCK SHEET CUTTING PLANS & NESTING BOM (GROUPED BY MATERIAL TYPE)"
    ws3["A1"].font = font_section_hdr

    n_row = 3
    to_m2 = 1.0 / 1_000_000.0 if unit == "mm" else 1.0

    if material_nest_dict:
        for mat_name, n_res in sorted(material_nest_dict.items()):
            # Material Header Banner
            ws3.cell(row=n_row, column=1, value=f"MATERIAL: {mat_name.upper()} — NESTING CUT PLAN ({n_res.total_sheets} SHEETS)").font = font_mat_hdr
            ws3.merge_cells(start_row=n_row, start_column=1, end_row=n_row, end_column=9)
            ws3.cell(row=n_row, column=1).fill = fill_blue_header
            ws3.cell(row=n_row, column=1).alignment = align_left
            n_row += 1

            # BOM Headers
            bom_hdrs = [
                "Stock Sheet #", f"Stock Size ({unit})", "Pieces Placed",
                f"Used Area ({'m²' if unit=='mm' else 'sq.ft'})", f"Sheet Area ({'m²' if unit=='mm' else 'sq.ft'})",
                "Yield %", "Waste %", "Items Nested on Sheet", "Origin Pages"
            ]
            for c_idx, h in enumerate(bom_hdrs, start=1):
                cell = ws3.cell(row=n_row, column=c_idx, value=h)
                cell.font = font_table_hdr
                cell.fill = fill_slate_header
                cell.alignment = align_center

            start_bom_row = n_row + 1
            n_row += 1

            for s in n_res.sheets:
                items_str = ", ".join([p.tag for p in s.pieces])
                pages_str = ", ".join(sorted(list(set(f"P.{p.page_number}" for p in s.pieces))))

                ws3.cell(row=n_row, column=1, value=f"Sheet #{s.sheet_index}").font = font_bold
                ws3.cell(row=n_row, column=2, value=f"{int(s.stock_width)} × {int(s.stock_height)}").alignment = align_center
                ws3.cell(row=n_row, column=3, value=len(s.pieces)).alignment = align_center
                ws3.cell(row=n_row, column=4, value=round(s.used_area * to_m2, 3)).number_format = "#,##0.000"
                ws3.cell(row=n_row, column=5, value=round(s.sheet_area * to_m2, 3)).number_format = "#,##0.000"
                ws3.cell(row=n_row, column=6, value=s.yield_pct / 100.0).number_format = "0.0%"
                ws3.cell(row=n_row, column=7, value=s.waste_pct / 100.0).number_format = "0.0%"
                ws3.cell(row=n_row, column=8, value=items_str).font = font_normal
                ws3.cell(row=n_row, column=9, value=pages_str).alignment = align_center

                for c_idx in range(1, 10):
                    c_cell = ws3.cell(row=n_row, column=c_idx)
                    c_cell.border = cell_border
                    if n_row % 2 == 1:
                        c_cell.fill = fill_zebra
                n_row += 1

            # Subtotal for this Material
            ws3.cell(row=n_row, column=1, value=f"SUBTOTAL: {mat_name}").font = font_bold
            ws3.cell(row=n_row, column=2, value="-").alignment = align_center
            ws3.cell(row=n_row, column=3, value=f"=SUM(C{start_bom_row}:C{n_row-1})").font = font_bold
            ws3.cell(row=n_row, column=4, value=f"=SUM(D{start_bom_row}:D{n_row-1})").font = font_bold
            ws3.cell(row=n_row, column=4).number_format = "#,##0.000"
            ws3.cell(row=n_row, column=5, value=f"=SUM(E{start_bom_row}:E{n_row-1})").font = font_bold
            ws3.cell(row=n_row, column=5).number_format = "#,##0.000"
            ws3.cell(row=n_row, column=6, value=f"=D{n_row}/E{n_row}").font = font_bold
            ws3.cell(row=n_row, column=6).number_format = "0.0%"
            ws3.cell(row=n_row, column=7, value=f"=1-F{n_row}").font = font_bold
            ws3.cell(row=n_row, column=7).number_format = "0.0%"
            ws3.cell(row=n_row, column=8, value="-").alignment = align_center
            ws3.cell(row=n_row, column=9, value="-").alignment = align_center

            for c_idx in range(1, 10):
                c_cell = ws3.cell(row=n_row, column=c_idx)
                c_cell.border = total_border
                c_cell.fill = fill_total_row
            n_row += 2

            # CNC Cut Coordinates Table for this Material
            ws3.cell(row=n_row, column=1, value=f"CNC SHOP CUTTING COORDINATES — {mat_name.upper()}").font = font_section_hdr
            n_row += 1

            cut_headers = ["Sheet #", "Panel Tag", "Page No.", "Room Name", f"Cut X ({unit})", f"Cut Y ({unit})", f"Length ({unit})", f"Width ({unit})", "Rotated 90°"]
            for c_idx, h in enumerate(cut_headers, start=1):
                cell = ws3.cell(row=n_row, column=c_idx, value=h)
                cell.font = font_table_hdr
                cell.fill = fill_slate_header
                cell.alignment = align_center
            n_row += 1

            for s in n_res.sheets:
                for p in s.pieces:
                    ws3.cell(row=n_row, column=1, value=f"Sheet #{s.sheet_index}").alignment = align_center
                    ws3.cell(row=n_row, column=2, value=p.tag).font = font_bold
                    ws3.cell(row=n_row, column=3, value=f"Page {p.page_number}").alignment = align_center
                    ws3.cell(row=n_row, column=4, value=p.sheet_name).font = font_normal
                    ws3.cell(row=n_row, column=5, value=round(p.x, 1)).number_format = "#,##0.0"
                    ws3.cell(row=n_row, column=6, value=round(p.y, 1)).number_format = "#,##0.0"
                    ws3.cell(row=n_row, column=7, value=round(p.width, 1)).number_format = "#,##0.0"
                    ws3.cell(row=n_row, column=8, value=round(p.height, 1)).number_format = "#,##0.0"
                    ws3.cell(row=n_row, column=9, value="YES" if p.is_rotated else "NO").alignment = align_center

                    for c_idx in range(1, 10):
                        c_cell = ws3.cell(row=n_row, column=c_idx)
                        c_cell.border = cell_border
                        if n_row % 2 == 1:
                            c_cell.fill = fill_zebra
                    n_row += 1
            n_row += 3
    else:
        ws3.cell(row=4, column=1, value="No nesting optimization results generated.").font = font_normal

    # =========================================================================
    # SHEET 4: QUALITY CONTROL & DRAFTING AUDIT LOG (IF PROVIDED)
    # =========================================================================
    if qc_report is not None and hasattr(qc_report, "issues") and qc_report.issues:
        ws4 = wb.create_sheet(title="QC Audit Log")
        ws4.views.sheetView[0].showGridLines = True

        ws4["A1"] = "CONTRACTOR DRAFTING QUALITY CONTROL & GEOMETRY AUDIT"
        ws4["A1"].font = font_section_hdr

        qc_hdrs = ["Issue ID", "Severity", "Category", "Element Tag", "Title", "Technical Description", "Remedy / Correction"]
        for c_idx, h in enumerate(qc_hdrs, start=1):
            cell = ws4.cell(row=3, column=c_idx, value=h)
            cell.font = font_table_hdr
            cell.fill = fill_amber_header
            cell.alignment = align_center

        qc_row = 4
        for issue in qc_report.issues:
            ws4.cell(row=qc_row, column=1, value=issue.id).font = font_bold
            ws4.cell(row=qc_row, column=2, value=issue.severity).alignment = align_center
            ws4.cell(row=qc_row, column=3, value=issue.category).font = font_normal
            ws4.cell(row=qc_row, column=4, value=issue.element_tag or "-").alignment = align_center
            ws4.cell(row=qc_row, column=5, value=issue.title).font = font_bold
            ws4.cell(row=qc_row, column=6, value=issue.description).font = font_normal
            ws4.cell(row=qc_row, column=7, value=issue.remedy).font = font_normal

            for c_idx in range(1, 8):
                c_cell = ws4.cell(row=qc_row, column=c_idx)
                c_cell.border = cell_border
                if qc_row % 2 == 1:
                    c_cell.fill = fill_zebra
            qc_row += 1

    # If api_usage_logs is None, check streamlit session_state
    if api_usage_logs is None:
        try:
            import streamlit as _st
            if hasattr(_st, "session_state") and "api_usage_logs" in _st.session_state:
                api_usage_logs = _st.session_state.api_usage_logs
        except Exception:
            pass

    # =========================================================================
    # SHEET 5: API USAGE & CLOUD COMPUTING COST REPORT
    # =========================================================================
    add_api_cost_report_sheet(wb, api_usage_logs=api_usage_logs, project_name=project_name)

    # Auto-fit Column Widths across all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    if "\n" in val_str:
                        val_str = max(val_str.split("\n"), key=len)
                    if not str(cell.value).startswith("="):
                        if len(val_str) <= 60:
                            max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3.5, 12.0), 55.0)

    # Save to buffer / file
    buffer = io.BytesIO()
    wb.save(buffer)
    xlsx_bytes = buffer.getvalue()

    if output_path:
        out_p = Path(output_path)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        with open(out_p, "wb") as f:
            f.write(xlsx_bytes)

    return xlsx_bytes
